from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ICUMS_Simulator_User_Guide.docx"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def borders(table, color="D9D9D9"):
    tblPr = table._tbl.tblPr
    b = tblPr.first_child_found_in("w:tblBorders")
    if b is None:
        b = OxmlElement("w:tblBorders")
        tblPr.append(b)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        tag = "w:" + edge
        el = b.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            b.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), color)

def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    borders(t)
    for i, h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c, "1F4E79")
        for r in c.paragraphs[0].runs:
            r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    for ridx, row in enumerate(rows):
        cells=t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text=str(val)
            cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2: shade(cells[i], "EEF4F8")
    doc.add_paragraph()
    return t

def h(doc, text, level=1):
    p=doc.add_heading(text, level=level)
    p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(6)
    return p

def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")

def app(doc, name, purpose, expect, roles):
    h(doc, name, 2)
    doc.add_paragraph(purpose)
    doc.add_paragraph("What to expect", style="Heading 3")
    for x in expect: bullet(doc, x)
    doc.add_paragraph("Who uses it", style="Heading 3")
    add_table(doc, ["Role", "Typical use"], roles)

doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles=doc.styles
styles["Normal"].font.name="Aptos"; styles["Normal"].font.size=Pt(10); styles["Normal"].paragraph_format.space_after=Pt(6)
for s in ("Title","Heading 1","Heading 2","Heading 3"):
    styles[s].font.name="Aptos Display"; styles[s].font.color.rgb=RGBColor(0,0,0)
styles["Title"].font.size=Pt(28); styles["Heading 1"].font.size=Pt(18); styles["Heading 2"].font.size=Pt(14); styles["Heading 3"].font.size=Pt(11)

p=doc.add_paragraph(); p.style="Title"; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("ICUMS Simulator User Guide")
p=doc.add_paragraph("Student, Instructor, and Administrator Handbook"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.size=Pt(15); p.runs[0].font.color.rgb=RGBColor(31,78,121)
p=doc.add_paragraph("Training simulator only - not official ICUMS"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.bold=True

doc.add_paragraph("This guide explains what each major application area does, who normally uses it, and what should happen when you use it. It is written for the current local-first implementation and uses fictional training content. The simulator does not submit customs declarations, process payments, release cargo, or create government records.")

h(doc,"How to read this guide")
add_table(doc,["If you are...","Start here"],[
("A student","Your journey, then Accounts, Authentication, Learning, Assessments, Scenarios, Evaluations, and Reports."),
("An instructor","Instructor responsibilities, then Onboarding, Progress, Evaluations, Reports, and Audit."),
("An administrator","Admin responsibilities, then Accounts, Authentication and Authorization, Onboarding, Learning, Scenarios, Reports, Audit, and backups.")])

h(doc,"The overall journey")
doc.add_paragraph("The application is designed as one connected learning path:")
for x in ["Create or receive an account and sign in.","Accept the current simulator disclaimer.","Complete the theory lessons and module assessment.","Pass the untimed final theory examination.","Complete simulator orientation.","Practise the guided fictional Import scenario.","Complete the independent competency scenario when assigned.","Review practical evidence and, when all completion rules are met, open the training record and certificate."]: bullet(doc,x)
doc.add_paragraph("A guided practice result does not issue a certificate. Completion requires orientation, a passed final theory attempt, and a passed competency-purpose practical evaluation.")

h(doc,"Important terms")
add_table(doc,["Term","Meaning"],[
("Training simulator","A safe learning environment using fictional records and scenarios."),
("Training disclaimer","The notice that the simulator is not official ICUMS and does not create government records."),
("Enrolment","The link between a student account and a published programme version."),
("Scenario attempt","One saved run through a practical workflow."),
("System outcome","The automatic practical pass/fail calculated from the frozen rubric and action history."),
("Effective outcome","The outcome currently used after any permitted instructor revision."),
("Completion record","The training record created after the qualification rules are satisfied."),
("Certificate","A simulator-training certificate with a unique number; it is not an official customs qualification.")])

app(doc,"Accounts","Stores user profiles, roles, and the relationships that identify who is learning or administering the simulator.",[
"Users have a UUID identity, username, email, and password.",
"One person can hold more than one role, such as Instructor and Administrator.",
"Administrators normally create users and enrol students through Django Admin."],[
("Student","Uses their own account and sees their own learning and records."),
("Instructor","Uses account identity when reviewing students and giving feedback."),
("Administrator","Creates users, assigns groups, and manages account access.")])

app(doc,"Authentication and Authorization","Authentication proves who is signing in. Authorization decides which pages and actions that signed-in user may use.",[
"Students can access their own dashboard, lessons, assessments, practical attempts, results, and certificates.",
"Instructors can review student progress, practical evaluations, feedback, and approval actions where enabled.",
"Administrators can manage trusted configuration and content in Django Admin.",
"Every authenticated training page carries the simulator warning."],[
("Student","Sign in, sign out, accept the disclaimer, and use permitted student pages."),
("Instructor","Sign in and access instructor-only review features."),
("Administrator","Sign in to the admin area and manage role groups.")])

app(doc,"Onboarding","Connects a user to a programme and records the simulator disclaimer acceptance.",[
"A published programme version is selected for the enrolment.",
"An active enrolment is required for the student journey.",
"The current disclaimer must be accepted before training pages are available.",
"The enrolling administrator is recorded for traceability."],[
("Student","Accepts the current disclaimer and confirms access to the assigned programme."),
("Instructor","Checks that a student has an active enrolment and accepted disclaimer."),
("Administrator","Creates the user, adds the Student group automatically through enrolment, and selects the programme version.")])

app(doc,"Learning","Delivers ordered theory modules, lessons, content blocks, resources, and knowledge checks.",[
"Lessons save completion and the learner's last position.",
"Published modules and lessons appear in the roadmap.",
"Resources can be linked to weak areas for remediation.",
"Fictional content is clearly labelled and replaceable when approved content is supplied."],[
("Student","Reads lessons, completes blocks, answers embedded checks, and follows remediation resources."),
("Instructor","Reviews whether students have completed modules and where they are stuck."),
("Administrator","Creates, publishes, and maintains learning content through the admin area.")])

app(doc,"Assessments","Runs untimed module and final theory assessments with frozen question sets and recorded attempts.",[
"The final theory assessment remains locked until published module lessons are complete.",
"Time may be recorded for insight, but it does not affect pass/fail.",
"Each attempt is numbered and retained; earlier attempts are not overwritten.",
"Question selections and responses are preserved for later review."],[
("Student","Starts an attempt, answers questions, submits, and reviews the result and remediation."),
("Instructor","Reviews attempts, scores, and student progress."),
("Administrator","Publishes assessment definitions, questions, pass thresholds, and attempt policies.")])

app(doc,"Progress","Combines learning and programme milestones into a clear view of what is complete and what is unlocked.",[
"Module completion contributes to final-theory access.",
"Passing final theory unlocks orientation.",
"Completing orientation unlocks practical training.",
"The dashboard shows the current state without hiding previous evidence."],[
("Student","Uses the dashboard and roadmap to decide what to do next."),
("Instructor","Uses progress views to identify students needing support."),
("Administrator","Checks that programme structure and unlock rules behave as configured.")])

app(doc,"Scenarios","Runs the practical simulator as a versioned state machine using fictional documents, actions, conditions, and effects.",[
"A scenario version is published and then used to create an attempt.",
"Available actions depend on the current state and stored conditions.",
"Attempts save automatically and can be resumed.",
"Parallel actions, such as fictional shipping-line work, can proceed beside the main workflow.",
"Every action is recorded with before and after state evidence.",
"Beginner assistance can show hints; competency mode does not show hints."],[
("Student","Reviews fictional documents, chooses permitted actions, saves progress, and completes the workflow."),
("Instructor","Reviews attempts and action history; uses scenario evidence during coaching."),
("Administrator","Creates and publishes scenario versions, actions, conditions, documents, and attempt policies.")])

app(doc,"Evaluations","Calculates a practical result from a frozen rubric version and the completed action history.",[
"Criteria can check completion, required actions, state flags, and assistance usage.",
"The system score and system outcome remain preserved.",
"Mandatory criterion failures can make a result fail even when the total score is high.",
"Instructors may add visible feedback or private notes.",
"A permitted instructor revision changes the effective outcome but does not erase the original system outcome."],[
("Student","Reads strengths, unmet criteria, feedback, remediation, and informational elapsed time."),
("Instructor","Reviews evidence, gives feedback, and revises an outcome only with a reason when authorized."),
("Administrator","Maintains rubric versions, criteria, thresholds, and remediation links.")])

app(doc,"Reports","Turns qualifying evidence into a completion record and simulator-training certificate.",[
"Completion is created only after orientation, final theory pass, and a passed competency practical.",
"Practice-only scenarios cannot issue a certificate.",
"The completion record stores evidence and policy snapshots for auditability.",
"Certificate issuance is idempotent, so retrying does not create duplicates.",
"The certificate clearly states that it is simulator training and not an official qualification."],[
("Student","Opens the training record and prints the certificate when eligible."),
("Instructor","Approves pending completion when the policy requires it."),
("Administrator","Configures completion policy and reviews completion or certificate records.")])

app(doc,"Audit","Keeps append-only evidence of important security, enrolment, disclaimer, evaluation, and approval events.",[
"Login, logout, disclaimer acceptance, enrolment, evaluation revisions, and completion approvals can be traced.",
"Audit entries identify the actor, action, target, summary, and time.",
"Audit records are evidence for review, not a replacement for the student-facing workflow."],[
("Student","Normally does not edit or manage audit records."),
("Instructor","Uses audit evidence when reviewing decisions and approvals."),
("Administrator","Reviews audit events and protects the audit history.")])

app(doc,"Django Admin","Provides the trusted administrative workspace for configuration and initial content loading.",[
"Admin users manage accounts, roles, programmes, theory, scenarios, rubrics, completion policy, and read-only evidence.",
"Published records have safeguards against casual editing where changing them would invalidate history.",
"The admin area is not the student interface and should not be exposed casually to the public internet."],[
("Student","Does not use Django Admin."),
("Instructor","May receive limited admin access only if assigned by an administrator."),
("Administrator","Uses it for account creation, enrolment, content, policy, and operational review.")])

h(doc,"Role based quick reference")
add_table(doc,["Role","First tasks","Main areas","What success looks like"],[
("Student","Sign in; accept disclaimer; open dashboard","Learning, Assessments, Progress, Scenarios, Evaluations, Reports","Can explain the lesson, complete required assessments, and see evidence for results."),
("Instructor","Open progress; review attempts; support learners","Progress, Evaluations, Reports, Audit","Can give traceable feedback and approve completion only when policy allows."),
("Administrator","Create account; enrol; publish content; configure policy","Accounts, Onboarding, Learning, Scenarios, Reports, Audit, Admin","The right user sees the right programme and every important decision is traceable.")])

h(doc,"Common questions")
for q,a in [
("Why can I see a practice result but no certificate?","Practice is for guided learning. A certificate requires a passed competency-purpose scenario as well as final theory and orientation."),
("Why is an action unavailable?","The action is not valid for the current state or its conditions have not been met. Review the current documents, status, and parallel work."),
("Does time affect my score?","No. Time is recorded as informational evidence only."),
("Can an instructor change a result?","Only when authorized. The original automatic result remains visible and the revision requires a reason and is audited."),
("Is this the real ICUMS system?","No. It is a training simulator using fictional documents and transactions. Nothing is submitted to government systems.")]:
    p=doc.add_paragraph(); p.add_run(q+" ").bold=True; p.add_run(a)

h(doc,"Where to get help")
doc.add_paragraph("Students should contact their instructor when a lesson, assessment, or scenario is unclear. Instructors should contact the administrator when access, enrolment, content, or policy is incorrect. Administrators should preserve the audit trail and record any change to programme, rubric, or certificate policy.")
doc.add_paragraph("This guide describes the current fictional implementation. Customs rules, practical criteria, official terminology, and visual alignment with real ICUMS screens remain subject to approved reference material and subject-matter review.")

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run("ICUMS Simulator User Guide | Training simulator only | Not official ICUMS")
doc.save(OUT)
print(OUT)

