from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"docs"/"ICUMS_Simulator_Detailed_Screen_Reference.docx"
def shade(c, fill):
    shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); c._tc.get_or_add_tcPr().append(shd)
def table(doc, headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,x in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=x; shade(c,"1F4E79")
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    for n,row in enumerate(rows):
        cs=t.add_row().cells
        for i,x in enumerate(row): cs[i].text=str(x)
        if n%2==1:
            for c in cs: shade(c,"EEF4F8")
    doc.add_paragraph()
def h(doc,text,l=1): doc.add_heading(text,l)
def bullets(doc,items):
    for x in items: doc.add_paragraph(x,style="List Bullet")
def screen(doc,title,route,who,what,fields,expect):
    h(doc,title,2); doc.add_paragraph("Route or entry point: "+route); doc.add_paragraph("Who uses it: "+who); doc.add_paragraph(what)
    doc.add_paragraph("Fields and controls",style="Heading 3"); table(doc,["Item","What to enter or select","What happens"],fields)
    doc.add_paragraph("What to expect",style="Heading 3"); bullets(doc,expect)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
for s in ("Title","Heading 1","Heading 2","Heading 3"):
    doc.styles[s].font.name="Aptos"; doc.styles[s].font.color.rgb=RGBColor(0,0,0)
doc.styles["Normal"].font.name="Aptos"; doc.styles["Normal"].font.size=Pt(9.5)
p=doc.add_paragraph(); p.style="Title"; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("ICUMS Simulator Detailed Screen Reference")
p=doc.add_paragraph("What appears on each screen and what each field means"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph("Training simulator only - not official ICUMS"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.bold=True

doc.add_paragraph("This is a companion to the role-based user guide. It follows the current files and templates in the project, so labels may change as the interface develops. It describes the simulator's fictional training data, not official customs procedures.")
h(doc,"How to use this reference")
doc.add_paragraph("For every screen, first check the route or menu entry, then identify the fields, then read the expected result. If a field is not shown in the current browser screen, it is usually an administrator-only model field or a future extension.")
table(doc,["Area","Typical screens covered","Primary role"],[
("Accounts and Admin","Users list, user detail, enrolment inline, groups","Administrator"),
("Onboarding","Disclaimer page, programme and enrolment records","Student, Administrator"),
("Learning and Progress","Roadmap, lesson, dashboard, orientation","Student, Instructor"),
("Assessments","Assessment take, result, attempt evidence","Student, Instructor"),
("Scenarios","Scenario list, detail, workspace, actions, hints","Student, Instructor"),
("Evaluations","Evaluation result, feedback, revision","Student, Instructor"),
("Reports","Completion record, certificate, policy","All roles"),
("Audit","Audit event list and detail","Instructor, Administrator")])

h(doc,"Accounts and administrator screens")
screen(doc,"Users list","Admin > Accounts > Users","Administrator","This page lists user accounts. Select a username to open that user's detail page.",[
("Search","Username or email (where enabled)","Narrows the list to matching accounts."),
("Username","Account login name","Opens the user detail form."),
("Email","Contact address","Used for identity and future communications; it is not a password."),
("Active / staff / superuser","Account status flags","Controls whether the account can sign in and whether it can use admin features; assign carefully.")],["The list is not the student learning dashboard.","Use the username link to manage the account.","Never share passwords or grant superuser access casually."])
screen(doc,"User detail","Admin > Accounts > Users > select a username","Administrator","The user detail page combines account credentials, personal information, permissions, and enrolments.",[
("Username","Unique login name","Identifies the account and is used at sign-in."),
("Password","Set or change through the password control","Passwords are stored securely; the existing password is not displayed in readable form."),
("First name / Last name","Student or staff name","Appears in personalisation and certificates where available."),
("Email address","User's email","Stores contact information."),
("Active","Checked for a usable account","An inactive user cannot use normal sign-in."),
("Staff status","Checked only for admin users","Allows access to Django Admin when permissions also allow it."),
("Superuser status","Use only for trusted system owners","Grants all admin permissions; normally leave unchecked."),
("Groups","Student, Instructor, Administrator","Assigns role capabilities. A user may have multiple groups."),
("Enrolments inline","Programme version and status","Links the user to a published programme. Saving an enrolment automatically assigns Student and records the enrolling administrator.")],["The password field is for setting a password, not reading one.","Role membership and enrolment are separate: a user may be Instructor without being enrolled as a student.","Use Save and continue editing if you want to confirm the inline enrolment."])
screen(doc,"Enrolment inline","Inside a user detail page","Administrator","Use this section to enrol a student without a command line.",[
("Student","Pre-filled with the current user","The account receiving the enrolment."),
("Programme version","Published programme version","Determines the modules, assessments, and completion policy available."),
("Status","Active, suspended, or completed","Active allows the normal learning journey; other statuses restrict access or describe history."),
("Enrolled by","Recorded automatically","Preserves which administrator created the relationship.")],["A student needs an active enrolment to see the roadmap and practical training.","The Student group is assigned automatically when an enrolment is saved.","Do not delete an enrolment to erase history; retain it and change status when appropriate."])

h(doc,"Authentication and onboarding")
screen(doc,"Login","/accounts/login/","All users","Enter credentials for the account created by an administrator.",[
("Username","Account username","Identifies the account."),
("Password","Account password","Checks the password securely."),
("Sign in","Submit button","Starts an authenticated session.")],["Invalid credentials return a safe error without revealing which value was wrong.","After sign-in, the system may redirect to the disclaimer before the dashboard.","Use Sign out when leaving a shared computer."])
screen(doc,"Simulator disclaimer","/onboarding/disclaimer/","Student and any authenticated user who has not accepted the current version","Read the training notice and accept it before continuing.",[
("Disclaimer text","Read-only current notice","Explains that documents and transactions are fictional and not official ICUMS."),
("Accept","Confirmation button","Records acceptance against the current disclaimer version and creates an audit event.")],["A new disclaimer version can require acceptance again.","Until accepted, protected training pages redirect here."])

h(doc,"Learning and progress")
screen(doc,"Student dashboard","/","Student","The dashboard is the starting point after the disclaimer.",[
("Enrolment banner","Current programme and status","Confirms which programme version is active."),
("Module cards","Module title, progress, next lesson","Links to the roadmap or lesson."),
("Final theory","Locked or available status","Shows whether module prerequisites are complete."),
("Orientation","Locked or complete status","Becomes available after a passed final theory assessment."),
("Practical","Locked or available status","Becomes available after orientation."),
("Training record","Completion status and certificate link","Appears when a qualifying completion record exists.")],["The dashboard is a progress summary, not an admin configuration screen.","Locked actions explain which prerequisite is missing."])
screen(doc,"Roadmap and lesson","/learning/ and /learning/<module>/<lesson>/","Student","The roadmap lists published modules. A lesson displays ordered content blocks and any embedded knowledge check.",[
("Module card","Title and completion state","Opens the first available lesson."),
("Content block","Notice, text, checklist, or other published block","Marks the learner's place as they progress."),
("Knowledge check","Select an answer and submit","Shows immediate correctness and explanation; responses remain recorded."),
("Complete lesson","Completion control","Records lesson progress and unlocks the next lesson when prerequisites are met."),
("Resource link","Remediation or reference resource","Returns the learner to targeted support content.")],["Leaving and returning restores the saved position.","Unpublished content is not shown to students.","Fictional examples are labelled as training-only."])
screen(doc,"Orientation","/orientation/","Student","Orientation confirms that the student understands the simulator workspace before practical access.",[
("Orientation summary","Read-only instructions","Explains the practical workspace and simulator boundaries."),
("Complete orientation","Confirmation button","Records orientation_completed_at and an audit event.")],["Orientation is available only after final theory is passed.","Completing it unlocks Scenarios."])

h(doc,"Assessments")
screen(doc,"Assessment take","/assessments/<assessment-id>/","Student","Start or resume an untimed module or final-theory attempt.",[
("Question prompt","Read-only frozen question","The question set is frozen when the attempt begins."),
("Answer options","Select one available answer","Stores the selected response for this attempt."),
("Submit","Submit all answers","Calculates score and pass/fail; the attempt becomes read-only."),
("Attempt number","Displayed or implied","Each retry creates a new numbered attempt; older evidence remains.")],["The final theory assessment stays locked until published module lessons are complete.","Elapsed time may be recorded but never changes the score.","Incorrect answers may link to remediation resources."])
screen(doc,"Assessment result","/assessments/attempts/<attempt-id>/","Student and Instructor","Review the submitted score, outcome, responses, and remediation.",[
("Score and percentage","Calculated result","Shows points earned and the percentage."),
("Outcome","Pass or fail","Determines whether the next programme milestone is unlocked."),
("Response review","Selected answer and explanation","Explains correct and incorrect responses."),
("Remediation","Linked lesson or resource","Provides a targeted next step after a weak result.")],["A failed attempt is retained and does not block a permitted retry.","A passed final theory attempt unlocks orientation."])

h(doc,"Scenarios")
screen(doc,"Scenario list and detail","/scenarios/ and /scenarios/<version-id>/","Student","The list shows published practical scenarios available after orientation.",[
("Scenario title","Published scenario name","Opens the briefing page."),
("Purpose","Practice, module assessment, or competency","Explains whether the scenario is practice or can contribute to completion."),
("Assistance mode","Beginner, Intermediate, Advanced, or Competency","Explains whether hints or guidance are available."),
("Start or resume","Action button","Creates a new saved attempt or returns to an in-progress attempt.")],["The guided fictional Import scenario is practice-only.","The independent competency scenario does not provide hints.","Scenario rules are conceptual until approved reference material is supplied."])
screen(doc,"Scenario workspace","/scenarios/attempts/<attempt-id>/","Student and Instructor","The workspace is the practical transaction screen.",[
("Current state","Read-only workflow stage","Shows where the attempt currently is."),
("Fictional documents","Learner-visible fields and references","Provides safe training records; evaluator-only data is hidden."),
("Available actions","Buttons valid for the current state and conditions","Performs a transition and records before/after state evidence."),
("Parallel actions","Actions with no main state transition","Allows independent shipping-line work when conditions are met."),
("Hint","Beginner assistance control","Records an assistance event; unavailable in Advanced and Competency modes."),
("Action history","Sequence, action code, prior state, new state","Shows the append-only evidence trail.")],["The workspace saves after actions and can be resumed.","Unavailable actions are usually waiting for a prerequisite or the correct state.","A terminal state completes the attempt and triggers evaluation."])

h(doc,"Evaluations")
screen(doc,"Practical evaluation","/evaluations/<evaluation-id>/","Student and Instructor","The result page explains how the completed scenario was scored.",[
("System score","Points and percentage","Calculated from the published rubric version."),
("System outcome","Automatic pass or fail","Preserved even if an instructor later revises the effective result."),
("Effective outcome","Current pass or fail","The outcome used for downstream completion decisions."),
("Criterion results","Passed/unmet criteria, evidence, feedback","Explains strengths and weaknesses."),
("Instructor feedback","Visible or private notes","Visible notes help the student; private notes remain instructor-only."),
("Revision history","Original outcome, revised outcome, reason, actor, timestamp","Makes an override auditable.")],["Elapsed time is informational only.","Mandatory failures can fail an evaluation even when the percentage is high.","Only authorized instructors or administrators can revise outcomes."])

h(doc,"Reports and certificates")
screen(doc,"Completion record","/records/completion/<completion-id>/","Student, Instructor, Administrator","Shows whether the student has met the programme completion gate.",[
("Status","Pending instructor approval or Completed","Shows whether a certificate can be issued."),
("Programme","Programme name and version","Identifies the training completed."),
("Qualifying scenario","Competency scenario title","Confirms that the result came from a competency-purpose practical."),
("Practical result","Effective outcome and system score","Summarizes the qualifying evaluation."),
("Approve completion","Instructor-only action when policy requires it","Completes the record, records the approver, and issues the certificate.")],["Completion requires final theory pass, orientation, and passed competency practical.","Practice-only results do not create completion.","Evidence and policy snapshots are retained for auditability."])
screen(doc,"Certificate","/records/certificates/<certificate-id>/","Eligible Student, Instructor, Administrator","Displays a printable simulator-training certificate.",[
("Certificate number","Unique identifier","Can be used to distinguish the simulator record."),
("Student name","From the user profile","Uses full name where supplied, otherwise username."),
("Programme and date","From the completion record","States what simulator programme was completed and when."),
("Print","Browser print control","Opens the normal print dialog.")],["The certificate permanently states TRAINING SIMULATOR - NOT OFFICIAL ICUMS.","Issuance is idempotent: retrying does not create duplicates.","It is not a government-issued qualification."])

h(doc,"Audit and evidence")
screen(doc,"Audit events","Admin > Audit > Audit events","Instructor and Administrator","Audit is a read-only trail of important actions.",[
("Actor","User who performed the action","Identifies who signed in, enrolled, revised, or approved."),
("Action code","Machine-readable event type","Groups events such as login, disclaimer.accepted, enrolment.created, practical_evaluation.revised, and completion.approved."),
("Target","Object type and ID","Identifies what the event affected."),
("Summary","Human-readable description","Explains the event without replacing the underlying record."),
("Timestamp and IP","When and from where available","Supports review and incident investigation.")],["Audit entries are append-only.","Students normally do not edit or manage audit records.","Use the audit trail to explain a decision; do not use it to change a result."])

h(doc,"Administrator content and policy screens")
table(doc,["Admin area","What you configure","What to verify before publishing"],[
("Programmes and versions","Programme identity, version, status","Correct version is published and assigned to the intended students."),
("Modules and lessons","Order, titles, blocks, resources, checks","Prerequisites and fictional labels are clear."),
("Questions and assessments","Question versions, options, pass percentage, attempt limits","Question set is complete and final assessment is untimed."),
("Scenarios","Scenario version, purpose, assistance mode, states, actions, conditions, documents","Exactly one initial state, terminal path, and learner/evaluator data separation."),
("Rubrics and criteria","Points, mandatory flags, pass threshold, remediation","Criteria are evidence-based and still marked demonstration until approved."),
("Completion policy","Active flag, certificate template version, instructor approval switch","Only qualifying competency results can complete the programme."),
("Completion and certificates","Read-only evidence records","Never edit to correct history; use approved policy or revision workflows.")])

h(doc,"Troubleshooting by symptom")
table(doc,["Symptom","Likely reason","Next check"],[
("I am redirected to the disclaimer","Current disclaimer has not been accepted","Open Onboarding and accept the current notice."),
("The roadmap is empty","No active enrolment or no published modules","Ask an administrator to verify the user enrolment and programme version."),
("Final theory is locked","Published module lessons are incomplete","Review Progress and complete each required lesson."),
("Practical is locked","Orientation is incomplete","Pass final theory, open Orientation, and complete it."),
("An action is unavailable","State or condition is not satisfied","Read the current state, documents, and parallel work requirements."),
("No certificate appears","Completion prerequisites are not all satisfied or approval is pending","Check final theory, orientation, competency evaluation, and completion policy."),
("Admin page is unavailable","Account lacks staff/admin permission","Ask an administrator to review groups and staff status.")])

h(doc,"Support boundary")
doc.add_paragraph("This reference explains the current implementation. It does not replace approved customs procedures, official ICUMS training, or subject-matter review. If a screen or label conflicts with approved requirements, record the issue and ask the administrator or product owner before changing a published scenario, rubric, or certificate policy.")
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run("ICUMS Simulator Detailed Screen Reference | Training simulator only")
doc.save(OUT); print(OUT)

