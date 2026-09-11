from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "ICUMS_Simulator_Administrator_Manual.docx"
NAVY = "17324D"
PALE = "EEF3F6"
GRAY = "D9D9D9"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    tc_pr.append(element)


def border_cell(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), GRAY)
        borders.append(tag)


def set_cell_margin(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, NAVY)
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(8.5)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
            if r_idx % 2:
                shade(cells[i], PALE)
            for paragraph in cells[i].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            border_cell(cell)
            set_cell_margin(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)
    return paragraph


def numbered(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.add_run(text)
    return paragraph


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def fields(doc, title, rows):
    heading(doc, title, 3)
    add_table(doc, ["Field", "What to enter", "Behaviour and rules"], rows, [1.55, 2.45, 2.7])


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(28)
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 11)):
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
    styles[name].paragraph_format.keep_with_next = True

title = doc.add_paragraph(style="Title")
title.add_run("ICUMS Simulator Administrator Manual")
subtitle = doc.add_paragraph("Complete course creation field reference and publishing workflow")
subtitle.style = styles["Subtitle"]
doc.add_paragraph("Current system setup September 2026")
doc.add_paragraph("Training simulator only - not official ICUMS")
doc.add_paragraph()
intro = doc.add_paragraph()
intro.add_run("Purpose. ").bold = True
intro.add_run("This manual tells administrators and instructors exactly where to begin, what to enter in each content field, how conditional fields behave, and how to publish a complete course from theory through the final assessment and practical competency stage.")
doc.add_paragraph("The recommended starting point is Instructor workspace > Course Builder. Django Admin is the advanced workspace for detailed configuration, document production, enrolment, evidence review, and corrections that the guided builder does not yet expose.")

heading(doc, "Contents", 1)
for item in ["1 Roles and workspaces", "2 Complete course setup sequence", "3 Guided Course Builder fields", "4 Theory lessons and questions", "5 Theory assessments", "6 Practical simulator flows", "7 Fictitious shipping documents", "8 Practical rubrics and competency assessment", "9 Publishing and enrolment", "10 Learner flow and completion", "11 Advanced Django Admin reference", "12 Final quality checklist and troubleshooting"]:
    bullet(doc, item)

heading(doc, "Roles and workspaces", 1)
doc.add_paragraph("Use the guided Course Builder for normal course authoring. Use Django Admin when a relationship, version, rubric rule, document, enrolment, or evidence record needs advanced control.")
add_table(doc, ["Workspace", "Primary users", "Use it for"], [
    ("Instructor dashboard and Course Builder", "Instructor, Administrator", "Create draft courses, modules, theory lessons, lesson questions, module assessments, guided practical flows, review, and publish."),
    ("Django Admin", "Administrator and trusted advanced instructor", "Users, roles, enrolments, detailed question versions, final theory assessment, scenario state machines, rubrics, fictitious BL/invoice/packing documents, completion policy, and read-only evidence."),
    ("Student dashboard", "Student", "Theory roadmap, lesson reading, knowledge checks, assessments, orientation, practical scenarios, results, and certificates."),
    ("Simulator workspace", "Student", "Run the programmed practical flow. It saves attempts and actions without changing the course definition."),
    ("Practical portal", "Student", "Access Cargo, Clearance, Single Window, and scenario training areas. Return to Workspace goes back to the learning workspace."),
], [1.7, 1.35, 3.65])

heading(doc, "Complete course setup sequence", 1)
doc.add_paragraph("Follow this order. It prevents broken links, missing questions, inaccessible assessments, and courses that cannot be published.")
for step in [
    "Prepare the course outline: course name, module order, lesson outcomes, theory content, questions, practical steps, pass marks, and required learner documents.",
    "Open Instructor workspace > Course Builder. Choose Theory when the course contains reading and questions only, or Theory and Practical when it also includes simulated flows.",
    "Create the course. The system creates Programme version 1 in Draft status.",
    "Add modules in teaching order. Add each lesson, its reading blocks, knowledge-check question, answers, and explanation.",
    "Allow the lesson builder to add each lesson question to its module assessment, or configure assessment items manually in Django Admin.",
    "For a combined course, add one or more guided practicals to the relevant module. Enter the learner-facing steps in their correct sequence.",
    "Create the required fictitious Bill of Lading, Commercial Invoice, Proforma Invoice, and Packing List in Django Admin and connect each to the correct scenario version.",
    "Create or verify the final theory assessment in Django Admin. It belongs to the programme version, not to one module.",
    "For assessed practical work, confirm the scenario purpose, assistance mode, rubric, criteria, pass percentage, and mandatory rules.",
    "Open Review course. Resolve every issue. Publish only after checking content, answers, practical order, documents, and assessment settings.",
    "Create the student account and enrol the student into the published programme version. The Student role is added automatically when enrolment is created through the user page.",
    "Test with a student account: disclaimer, theory lessons, module assessment, final theory, orientation, practical, evaluation, and certificate rules.",
]: numbered(doc, step)

heading(doc, "Guided Course Builder fields", 1)
fields(doc, "Create course", [
    ("Course name", "A clear learner-facing name, such as International Trade and Customs", "Creates a new Programme and Draft Programme Version 1. A unique internal code is generated automatically."),
])
fields(doc, "Add module", [
    ("Module title", "One major topic or learning stage", "Used for navigation, assessment naming, and learner progress."),
    ("Description", "What learners should understand or be able to do", "Optional but recommended. It appears as module context."),
    ("Order", "Created automatically from the module sequence", "Advanced Admin can adjust it. Lower numbers appear first."),
    ("Prerequisite", "Earlier module that must be completed first", "Advanced Admin field. Leave blank when modules are independent."),
])
fields(doc, "Add theory lesson", [
    ("Lesson title", "Specific subject of the lesson", "Generates a unique lesson slug within the module."),
    ("Summary", "Short overview of what the lesson teaches", "Displayed before or beside lesson content."),
    ("Key concept heading", "Heading for the main reading section", "Creates the first Text content block."),
    ("Teaching content", "Complete theory explanation in learner-friendly language", "Required. This is reading material, not a simulator step."),
    ("Example heading and body", "A worked example or realistic fictional illustration", "Optional. Both belong to an Example content block."),
    ("Important note heading and body", "Warning, rule, exception, or reminder", "Optional. Creates a Notice content block."),
    ("Knowledge-check question", "One clear question testing the lesson", "The guided builder creates a Single Choice Question Version 1."),
    ("Correct answer", "The one correct option", "Saved as the first option and marked correct."),
    ("Incorrect answers", "Plausible but clearly incorrect alternatives", "Two are required and a third is optional. Do not use trick answers."),
    ("Explanation", "Why the correct answer is correct", "Shown after the learner answers. Explain the rule, not merely the letter or option."),
    ("Also add to module assessment", "Normally leave selected", "Creates or reuses the module assessment and adds this question to it."),
])
doc.add_paragraph("Lesson publication behaviour. Lessons, content blocks, question versions, module assessments, and practicals remain drafts while the course is being built. Publishing the course publishes the connected modules, lessons, lesson question versions, assessments, practical versions, and practical rubric versions.")

heading(doc, "Theory lessons and questions", 1)
doc.add_paragraph("Theory is reading with embedded knowledge checks and formal assessments. It does not contain simulator actions. Each lesson should have at least one teaching block and at least one knowledge check before publication.")
fields(doc, "Advanced lesson fields", [
    ("Module", "The draft module containing the lesson", "Must belong to the draft programme version."),
    ("Slug", "Short lowercase identifier", "Unique inside the module. The guided builder creates it automatically."),
    ("Order", "1, 2, 3 and so on", "Controls lesson sequence."),
    ("Is published", "Selected only when ready", "The guided course publish action sets it. Published programme content is protected from casual edits."),
    ("Content block kind", "Text, Example, or Notice", "Changes how the learner-facing section is presented."),
    ("After block", "The content block after which a check appears", "Must belong to the same lesson."),
])
fields(doc, "Question and answer fields", [
    ("Question code", "Stable unique identifier", "The conceptual question. Use a meaningful slug; do not reuse it for a different concept."),
    ("Version", "1 for the first wording", "Create a new version when changing a published question."),
    ("Question type", "Single choice or True or false", "The current learner UI expects one correct option."),
    ("Prompt", "The full question", "Avoid ambiguity and information not taught in the lesson."),
    ("Explanation", "Correct-answer rationale", "Learner feedback and remediation context."),
    ("Points", "Positive whole number", "Contributes to assessment score."),
    ("Remediation resource", "Optional linked reading", "Recommended when an incorrect answer should direct the learner to supporting material."),
    ("Is published", "Select only after options are complete", "Published question versions cannot be casually edited or deleted."),
    ("Answer label", "Text displayed as an option", "Use one row per answer."),
    ("Is correct", "Select exactly one option", "Course review blocks publication if a lesson check does not have exactly one correct option."),
    ("Answer order", "1, 2, 3 and so on", "Controls display order unless assessment randomisation changes question order."),
])

heading(doc, "Theory assessments", 1)
fields(doc, "Assessment fields", [
    ("Programme version", "The draft course version", "All assessment items must support this course."),
    ("Module", "Select for a module assessment; leave blank for final theory", "A Final Theory Examination cannot belong to one module."),
    ("Title", "Learner-facing assessment name", "Use consistent names such as Module 2 Assessment or Final Theory Examination."),
    ("Assessment type", "Module assessment or Final theory examination", "Controls scope and unlocking behaviour."),
    ("Pass percentage", "Required percentage such as 70", "Used to calculate pass or fail."),
    ("Maximum attempts", "Whole number or blank", "Blank means unlimited attempts."),
    ("Randomize questions", "Select when desired", "Changes question order but does not alter the frozen question set for an attempt."),
    ("Is published", "Select after items and answers are checked", "The course publish action publishes connected assessments."),
    ("Assessment items", "Published or ready Question Versions", "Order controls the normal sequence. A question version can appear only once per assessment."),
])
doc.add_paragraph("Final theory flow. The final theory examination remains locked until all published module lessons are complete. Time is recorded for information but does not affect pass or fail. Attempts are numbered and retained rather than overwritten.")

heading(doc, "Practical simulator flows", 1)
doc.add_paragraph("A practical contains teaching context plus a programmed sequence of simulator actions. The Simulator Workspace is the learner's sandbox: their actions create attempt records and do not change the course definition or other learners' work.")
fields(doc, "Guided practical builder", [
    ("Practical title", "Name of the workflow being practised", "Creates a Scenario and Scenario Version 1."),
    ("Area", "Import, Export, Transit, or Warehouse", "Controls classification and learner navigation."),
    ("Briefing", "Fictitious assignment and starting situation", "Tell the learner what has happened, their role, and the expected result."),
    ("Learning objective", "Observable skill the learner should demonstrate", "Used to explain why the practical exists."),
    ("Assistance mode", "Beginner, Intermediate, or Advanced", "Beginner includes hints; Intermediate has limited guidance; Advanced has no hints. Competency is configured in Advanced Admin."),
    ("Maximum attempts", "Whole number or blank", "Blank means unlimited attempts."),
    ("Steps", "One simulated action per line, in execution order", "Requires 2 to 30 steps. The builder creates states, transitions, completion effects, feedback, hints, and a completion rubric."),
])
fields(doc, "Advanced scenario version", [
    ("Scenario", "Parent scenario identity", "The title and area live on the Scenario record."),
    ("Module", "Course module containing this practical", "Links the practical into the course."),
    ("Version", "1 initially", "Create a new version instead of rewriting published history."),
    ("Status", "Draft, Published, or Retired", "Learners see published versions only."),
    ("Purpose", "Practice, Practical module assessment, or Practical competency assessment", "Practice does not qualify for a certificate. Competency can."),
    ("Assistance mode", "Beginner, Intermediate, Advanced, or Competency", "Purpose Competency requires Competency assistance mode."),
    ("Reference status", "Conceptual, Verified, Adapted, or To be confirmed", "Records the authority level of the simulated flow."),
    ("Initial data", "Advanced JSON key/value state only", "Leave as {} unless actions require flags. Keys are checked by conditions and changed by effects."),
])
fields(doc, "States and actions", [
    ("State key", "Unique lowercase identifier", "Unique inside the scenario version."),
    ("State label", "Learner-facing stage name", "Shown as the current practical stage."),
    ("Guidance", "Instruction for this stage", "Shown according to the assistance mode."),
    ("Is initial", "Select on exactly one state", "The learner attempt starts here."),
    ("Is terminal", "Select on the completion state", "Reaching it completes the practical."),
    ("Action code and label", "Stable code plus learner-facing action", "Code supports evidence; label is what the learner clicks."),
    ("From state and To state", "Current and next state", "Both must belong to the same scenario version. Blank states support parallel actions."),
    ("Conditions", "Required JSON key/value pairs", "The action appears only when the attempt state data satisfies them."),
    ("Effects", "JSON key/value updates", "Applied after the action, enabling later actions or recording evidence."),
    ("Success feedback", "Confirmation shown after action", "Tell the learner what changed."),
    ("Beginner hint", "Next-step guidance", "Available only where assistance permits it; assistance use can be evaluated."),
])

heading(doc, "Fictitious shipping documents", 1)
doc.add_paragraph("Create structured learner documents in Django Admin under Scenarios. Connect every document to the same Scenario Version used by the practical. Draft documents are author-only; published documents are available to learners who have an attempt for that scenario.")
heading(doc, "Bill of Lading", 2)
doc.add_paragraph("Use Guided form for entry, Live document view for immediate authoring preview, and View generated PDF after saving. Changing among Carrier Grid, Ocean Transport, and Multimodal changes presentation while retaining the entered data.")
fields(doc, "Bill of Lading document fields", [
    ("Scenario version", "Practical receiving the BL", "Determines learner access."), ("Status", "Draft or Published", "Publish only after checking the generated PDF."), ("Title and reference", "Fictitious document name and unique BL number", "Reference must be unique."), ("Template", "Carrier Grid, Ocean Transport, or Multimodal", "Changes layout, not shipment data."), ("Original status", "For example Non-negotiable training copy", "Printed document classification."), ("Number of originals", "Usually 0 for simulator copies", "Informational only."), ("Carrier and carrier agent", "Fictitious shipping line and agent", "Do not use a real transaction identity."), ("Shipper, consignee, notify party", "Full fictitious party details", "Keep identities consistent across invoice and packing list."), ("Booking and shipper references", "Fictitious cross-references", "Use consistent scenario identifiers."), ("Vessel and voyage", "Fictitious vessel and voyage reference", "Printed in the route section."), ("Receipt, loading, discharge, delivery", "Route places in sequence", "Place of receipt/delivery may differ from ports."), ("Freight terms", "Prepaid, collect, or training wording", "BL carriage term, not invoice line pricing."), ("Shipper declared value", "Optional BL declaration", "Not the commercial invoice total."), ("Issue and shipped dates", "Fictitious but internally consistent dates", "Shipment date should not follow issue events illogically."),
])
fields(doc, "BL cargo item fields and behaviour", [
    ("Cargo type", "Vehicle, General merchandise, Household or personal effects, Machinery, or Other cargo", "Controls which item fields appear."), ("Vehicle", "Year, make, model, VIN/chassis, HS code, weight and measurement", "Goods Description is hidden; vehicle identity fields are required."), ("General or personal effects", "Goods Description plus packages and weights", "Vehicle-specific fields are hidden."), ("Machinery", "Description plus optional year, manufacturer, model, serial/chassis and HS code", "Supports machinery with or without serial identity."), ("Other cargo", "All available item fields", "Use for mixed or unusual cargo."), ("Container and seal", "Container identifier, seal, and type", "Printed in the container/seal column."), ("Marks and numbers", "Marks printed on packages or cargo", "This identifies packages; it is not the narrative goods description."), ("Package quantity and type", "Count and form, such as 20 cartons", "Add another cargo item for each distinct vehicle or cargo line."), ("Gross weight and measurement", "Numeric value plus unit", "Printed per item and totalled in the PDF."),
])
heading(doc, "Commercial Invoice Proforma Invoice and Packing List", 2)
doc.add_paragraph("One editor handles all three. The selected type changes the heading, explanation, visible fields, live-preview columns, and generated PDF.")
add_table(doc, ["Type", "Purpose", "Pricing behaviour"], [
    ("Commercial Invoice", "Final seller-issued commercial document", "Shows unit price, line amount, FOB, Freight, Insurance, and C&F/CFR total."),
    ("Proforma Invoice", "Quotation before purchase or shipment", "Shows provisional prices and totals and prints a quotation-only warning."),
    ("Packing List", "Physical shipment and package detail", "Hides and rejects unit price, amount, FOB, Freight, and Insurance; shows packages, pieces, net weight, and gross weight."),
], [1.55, 2.5, 2.65])
fields(doc, "Commercial document fields", [
    ("Scenario version and status", "Target practical and Draft/Published", "Published learner access requires a scenario attempt."), ("Document type", "Commercial, Proforma, or Packing", "Immediately changes wording, fields, live preview, and PDF structure."), ("Title, reference, date", "Fictitious identity and date", "Reference must be unique."), ("Exporter name/address/contact", "Fictitious seller details", "Printed in the header."), ("Consignee name/address", "Fictitious buyer/importer", "Keep consistent with the BL."), ("Currency", "USD, EUR, GHS, etc.", "Used for invoice and proforma prices."), ("Container reference", "Container or proposed shipment reference", "Optional but useful for matching documents."), ("Payment terms", "Payment or quotation validity terms", "For proforma, state validity and provisional nature."), ("FOB", "Value of goods at export point", "Invoice/proforma only. Defaults visually to line total when not entered."), ("Freight and Insurance", "Separate commercial costs", "Invoice/proforma only."), ("Description and quantity", "One row per product or cargo item", "Used by all document types."), ("Package count and pieces per package", "For example 4 cartons and 5 pcs/carton", "Emphasised on packing lists."), ("Net and gross weight", "Weight excluding/including packaging", "Shown separately on packing lists and totalled."), ("Dimensions and HS code", "Optional supporting item details", "Store useful detail even where the standard PDF is concise."), ("Unit price and amount", "Price per unit and line total", "Commercial/proforma only. Packing-list validation prevents prices."),
])

heading(doc, "Practical rubrics and competency assessment", 1)
doc.add_paragraph("The guided practical builder creates a demonstration completion rubric. A final competency practical needs deliberate Advanced Admin configuration and subject-matter review.")
fields(doc, "Rubric fields", [
    ("Rubric code and title", "Stable identifier and readable name", "Use a new version when scoring policy changes."), ("Scenario version", "Exactly one practical version", "Each practical rubric version belongs to one scenario version."), ("Status", "Draft, Published, or Retired", "Publishing the course publishes connected rubric versions."), ("Pass percentage", "Required practical percentage", "Mandatory criteria can still cause failure."), ("Is demonstration", "Selected for unapproved practice rules", "Clear this only after formal approval of competency policy."), ("Criterion title and description", "The behaviour being evaluated", "Make evidence observable."), ("Dimension", "Accuracy, review, procedure, decision, correction, assistance, or completion", "Groups the skill measured."), ("Maximum points", "Positive whole number", "Contributes to total practical score."), ("Mandatory", "Select when failure must fail the practical", "Overrides a high overall percentage when unmet."), ("Evaluation rule", "Completion, required_actions, state_flags, or assistance_limit JSON", "Advanced field. The rule reads saved attempt evidence."), ("Remediation resource or scenario", "Follow-up support", "Presented when the criterion is not met."),
])

heading(doc, "Publishing and enrolment", 1)
heading(doc, "Course review gate", 2)
doc.add_paragraph("Review course blocks publishing when any module has no lesson, any lesson has no teaching content, any lesson has no knowledge check, a lesson check lacks exactly one correct answer, or a practical lacks at least two states and an action.")
doc.add_paragraph("Before selecting Publish course, manually confirm the final theory assessment, learner documents, practical purpose, competency rubric, completion policy, and realistic cross-document consistency. These advanced checks are not all enforced by the guided review gate.")
fields(doc, "Create and enrol a student", [
    ("Username, password, name, email", "The student's account details", "Use the Users area in Django Admin."), ("Groups", "Student and any authorised additional role", "Creating an enrolment through the user page automatically adds Student."), ("Programme version", "The published course version", "Do not enrol into a draft course for normal learning."), ("Enrolment status", "Active, Completed, or Withdrawn", "Active grants the learner journey. Completion is normally system-driven."), ("Enrolled by", "Recorded administrator", "Filled automatically when the enrolment is created in the user page."),
])

heading(doc, "Learner flow and final completion", 1)
add_table(doc, ["Stage", "Unlock or completion rule", "Administrative evidence"], [
    ("Disclaimer", "Learner accepts the current disclaimer", "Disclaimer acceptance and audit event."),
    ("Theory lessons", "Learner reads published lessons and completes checks", "Lesson progress and lesson check responses."),
    ("Module assessment", "Available for its module according to course flow", "Frozen assessment attempt, responses, score, and outcome."),
    ("Final theory", "Unlocks after all published module lessons are complete", "Submitted passing final-theory attempt."),
    ("Orientation", "Unlocks after final theory pass", "Orientation completion record."),
    ("Practical", "Unlocks after orientation", "Scenario attempt, state data, action history, and assistance events."),
    ("Practical evaluation", "Generated from the frozen rubric and action evidence", "System score/outcome plus any audited instructor revision."),
    ("Certificate", "Requires orientation, passed final theory, and passed competency-purpose practical", "Completion record and certificate. Practice-only scenarios do not qualify."),
], [1.35, 2.7, 2.65])

heading(doc, "Advanced Django Admin reference", 1)
add_table(doc, ["Area", "Create or configure", "Treat as read only or protected evidence"], [
    ("Accounts", "Users and role groups", "Authentication history is audited."),
    ("Onboarding", "Programmes, draft programme versions, enrolments, disclaimer versions", "Disclaimer acceptances."),
    ("Learning", "Modules, lessons, content blocks, resources", "Published historical content should be versioned, not rewritten."),
    ("Assessments", "Questions, question versions, options, assessments and items", "Theory attempts, responses, lesson-check responses."),
    ("Scenarios", "Scenarios, versions, states, actions, structured documents", "Attempts and action history."),
    ("Evaluations", "Rubrics, versions, criteria, remediation", "Practical evaluation, criterion results, revisions and feedback."),
    ("Reports", "Completion policy", "Completion records and certificates."),
    ("Audit", "No normal content creation", "Append-only security and administrative events."),
], [1.25, 2.65, 2.8])
doc.add_paragraph("Versioning rule. Draft records may be edited. Once content, a question, assessment, practical, rubric, or programme version has been published and used, create a new version for substantive changes so earlier learner evidence remains meaningful.")

heading(doc, "Final quality checklist and troubleshooting", 1)
for item in [
    "The course remains Draft while being assembled.", "Every module contains at least one lesson.", "Every lesson contains teaching content and a knowledge check.", "Every knowledge check has exactly one correct answer and an explanation.", "Module assessments contain the intended question versions.", "The final theory assessment belongs to the programme version and has no module selected.", "Each guided practical has a clear briefing, objective, ordered steps, initial state, terminal state, and valid transitions.", "Competency scenarios use Competency purpose and Competency assistance mode.", "Each practical has an appropriate rubric and approved mandatory rules.", "BL, invoice/proforma, and packing-list party, reference, quantity, package, weight, and route details agree.", "Every structured document PDF has been opened and visually checked before publication.", "Course Review shows no blocking issues.", "A test student can complete the full journey in the expected order.",
]: bullet(doc, item)
heading(doc, "Common problems", 2)
add_table(doc, ["Problem", "Likely reason", "Action"], [
    ("Theory page returns 403", "User lacks an active enrolment, Student access, or the content is not published", "Check user groups, active enrolment, programme status, and lesson/assessment publication."),
    ("Course cannot publish", "Review gate found missing module, lesson, content, question, correct option, or practical flow", "Open Review course and resolve each listed item."),
    ("Final theory is locked", "Published lessons are incomplete", "Complete the theory roadmap using a student account; do not manually alter evidence."),
    ("Practical is locked", "Final theory or orientation is incomplete", "Complete the preceding learner stages."),
    ("Document PDF returns 404", "Legacy document reference was removed or URL targets the wrong record type", "Use the structured document's own View generated PDF link after saving."),
    ("Packing list shows price fields", "Old static assets are cached or wrong document type selected", "Refresh the admin page and select Packing List; save and regenerate."),
    ("Live preview is unavailable", "The record is on the legacy Scenario Document editor", "Use Bills of Lading or Commercial Documents under Scenarios."),
    ("No certificate after practical", "The practical purpose is Practice or competency evidence is incomplete", "Require a passed competency-purpose practical plus passed final theory and orientation."),
], [1.75, 2.45, 2.5])

heading(doc, "Safe administration principles", 1)
for item in ["Use fictitious identities, references, values, vehicles, and transactions in learner content.", "Never represent simulator output as an official customs, banking, shipping, or commercial document.", "Keep Django Admin restricted to trusted staff and do not expose a development server to the public internet.", "Review audit records before correcting access, enrolment, evaluation, or completion disputes.", "Preserve published history by creating new versions instead of modifying evidence already used by learners."]:
    bullet(doc, item)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("ICUMS Simulator Administrator Manual  |  Training simulator only").font.size = Pt(8)

doc.core_properties.title = "ICUMS Simulator Administrator Manual"
doc.core_properties.subject = "Course creation field reference and publishing workflow"
doc.core_properties.author = "ICUMS Simulator"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
