from pathlib import Path
from docx import Document
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics

ROOT=Path(__file__).resolve().parents[1]
src=ROOT/"docs"/"ICUMS_Simulator_Detailed_Screen_Reference.docx"
out=ROOT/"docs"/"ICUMS_Simulator_Detailed_Screen_Reference.pdf"
docx=Document(src)
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name="GuideTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=25, leading=30, alignment=TA_CENTER, textColor=colors.HexColor("#1F4E79"), spaceAfter=8))
styles.add(ParagraphStyle(name="GuideSub", parent=styles["Normal"], fontName="Helvetica", fontSize=13, leading=16, alignment=TA_CENTER, textColor=colors.HexColor("#1F4E79"), spaceAfter=6))
styles.add(ParagraphStyle(name="H1x", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=colors.black, spaceBefore=12, spaceAfter=7))
styles.add(ParagraphStyle(name="H2x", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=17, textColor=colors.HexColor("#1F4E79"), spaceBefore=10, spaceAfter=5))
styles.add(ParagraphStyle(name="H3x", parent=styles["Heading3"], fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=6, spaceAfter=3))
styles.add(ParagraphStyle(name="Bodyx", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.3, leading=13, spaceAfter=5))
styles.add(ParagraphStyle(name="Bulletx", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.2, leading=12.5, leftIndent=12, firstLineIndent=-7, bulletIndent=0, spaceAfter=2))
styles.add(ParagraphStyle(name="Cellx", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.6, leading=10))
styles.add(ParagraphStyle(name="CellHead", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=7.7, leading=10, textColor=colors.white))
def esc(s):
    return str(s).replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")
def footer(canvas, doc):
    canvas.saveState(); canvas.setFont("Helvetica",7.5); canvas.setFillColor(colors.HexColor("#555555"))
    canvas.drawCentredString(A4[0]/2, 12*mm, f"ICUMS Simulator User Guide | Training simulator only | Not official ICUMS | {doc.page}")
    canvas.restoreState()
story=[]
for block in docx.element.body.iterchildren():
    if block.tag.endswith("}p"):
        p=next((x for x in docx.paragraphs if x._p is block), None)
        if not p: continue
        text=p.text.strip()
        if not text: story.append(Spacer(1,3)); continue
        style=p.style.name if p.style else ""
        if style=="Title": story.append(Paragraph(esc(text),styles["GuideTitle"]))
        elif style.startswith("Heading 1"): story.append(Paragraph(esc(text),styles["H1x"]))
        elif style.startswith("Heading 2"): story.append(Paragraph(esc(text),styles["H2x"]))
        elif style.startswith("Heading 3"): story.append(Paragraph(esc(text),styles["H3x"]))
        elif style=="List Bullet": story.append(Paragraph("• "+esc(text),styles["Bulletx"]))
        else: story.append(Paragraph(esc(text),styles["Bodyx"]))
    elif block.tag.endswith("}tbl"):
        tbl=next((x for x in docx.tables if x._tbl is block), None)
        if not tbl: continue
        data=[]
        for ridx,row in enumerate(tbl.rows):
            data.append([Paragraph(esc(cell.text), styles["CellHead" if ridx==0 else "Cellx"]) for cell in row.cells])
        t=Table(data, repeatRows=1, hAlign="LEFT")
        cmds=[("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#D9D9D9")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1F4E79")),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)]
        for i in range(1,len(data)):
            if i%2==0: cmds.append(("BACKGROUND",(0,i),(-1,i),colors.HexColor("#EEF4F8")))
        t.setStyle(TableStyle(cmds)); story.extend([Spacer(1,3),t,Spacer(1,6)])
SimpleDocTemplate(str(out), pagesize=A4, rightMargin=16*mm, leftMargin=16*mm, topMargin=15*mm, bottomMargin=18*mm, title="ICUMS Simulator User Guide", author="ICUMS Simulator").build(story, onFirstPage=footer, onLaterPages=footer)
print(out)
